import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_submission_doc(
    student_name="Rusheel Reddy",
    reg_number="230953176",
    repo_url="https://github.com/RK-0111/Fgai-assignment",
    output_docx="RusheelReddy_230953176_FGAI_Assignment2.docx",
    output_pdf="RusheelReddy_230953176_FGAI_Assignment2.pdf"
):
    doc = Document()

    # Page setup: 1-inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Header / Institution Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("Foundations of Generative AI (ICT4403)\n")
    run_inst.font.name = "Calibri"
    run_inst.font.size = Pt(16)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(31, 78, 121)

    run_sub = title_p.add_run("Assignment 2: Autoregressive Text Generation System Submission")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph() # Spacing

    # Section 1: Submission Details Table
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. Submission & Repository Identification")
    r_h1.font.name = "Calibri"
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(31, 78, 121)

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    details = [
        ("Student Name", student_name),
        ("Registration Number", reg_number),
        ("Assignment Title", "Assignment 2: Comparative Autoregressive Text Generation"),
        ("System Title", "Autoregressive Next-Word Generation System using Simple RNN, Bidirectional RNN, LSTM, and GRU"),
        ("Dataset Name & Source Link", "The Adventures of Sherlock Holmes (Sir Arthur Conan Doyle)\nProject Gutenberg: https://www.gutenberg.org/ebooks/1661"),
        ("GitHub Repository Link", repo_url)
    ]

    for i, (label, val) in enumerate(details):
        row = table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]

        c0.width = Inches(2.2)
        c1.width = Inches(4.3)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(4)
        p0.paragraph_format.space_after = Pt(4)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10.5)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(4)
        p1.paragraph_format.space_after = Pt(4)
        r1 = p1.add_run(val)
        r1.font.size = Pt(10.5)

    doc.add_paragraph() # Spacing

    # Section 2: Executive Summary & System Description
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("2. Executive System Overview")
    r_h2.font.name = "Calibri"
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(31, 78, 121)

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.15
    r_desc = p_desc.add_run(
        "This project develops, trains, and empirically compares an autoregressive text generation system "
        "formulated for next-word prediction across four distinct recurrent neural network architectures: "
        "Simple RNN, Bidirectional RNN, LSTM, and GRU. All four models were trained under strictly comparable "
        "hyperparameter budgets (embedding dimension 128, hidden dimension 128, 2 recurrent layers, Adam optimizer, "
        "lr = 0.002, Cross-Entropy Loss, and sequence length 12) on a curated corpus of Sir Arthur Conan Doyle's "
        "'The Adventures of Sherlock Holmes' from Project Gutenberg."
    )
    r_desc.font.size = Pt(10.5)

    # Section 3: Quantitative Results Summary
    h3 = doc.add_paragraph()
    r_h3 = h3.add_run("3. Quantitative Performance Summary")
    r_h3.font.name = "Calibri"
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(31, 78, 121)

    res_table = doc.add_table(rows=5, cols=7)
    res_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    res_table.style = "Table Grid"

    headers = ["Model", "Parameters", "Train Loss", "Val Loss", "Perplexity", "Top-1 Acc", "Time (s)"]
    for j, h_text in enumerate(headers):
        cell = res_table.rows[0].cells[j]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9.5)

    results_data = [
        ("Simple RNN", "708,548", "4.5351", "5.1805", "177.78", "15.01%", "51.13s"),
        ("Bidirectional RNN", "692,164", "4.4007", "5.2302", "186.84", "15.32%", "86.17s"),
        ("LSTM", "906,692", "4.6602", "5.1230", "167.84", "15.17%", "94.69s"),
        ("GRU", "840,644", "4.4772", "5.0970", "163.53", "16.32%", "146.44s")
    ]

    for i, row_data in enumerate(results_data):
        row = res_table.rows[i + 1]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            p = cell.paragraphs[0]
            if j == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(9.5)

    doc.add_paragraph() # Spacing

    # Section 4: Key Findings and Conclusion
    h4 = doc.add_paragraph()
    r_h4 = h4.add_run("4. Key Findings & Best-Performing Architecture")
    r_h4.font.name = "Calibri"
    r_h4.font.size = Pt(13)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(31, 78, 121)

    bullets = [
        "Optimal Model Selection: The Gated Recurrent Unit (GRU) achieved the lowest validation loss (5.0970), lowest validation perplexity (163.53), and highest top-1 accuracy (16.32%), providing superior autoregressive modeling quality.",
        "Computational Trade-offs: Simple RNN achieved the fastest convergence (51.13s) due to lack of gating logic, but suffered from higher perplexity. GRU balanced gating sophistication with fewer parameters (840k) compared to LSTM (906k).",
        "Temperature Sampling Dynamics: At T = 0.3, generations were deterministic with n-gram repetition loops. At T = 0.7, all models attained optimal lexical diversity (Distinct-1 ~0.77) and syntactical coherence. At T = 1.2, high entropy introduced grammatical degradation.",
        "Deliverables Provided: The full GitHub repository includes source code (.py), an interactive Jupyter Notebook (.ipynb), high-resolution metric figures, generated text JSON dumps, and a comprehensive README."
    ]

    for b in bullets:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(3)
        bp.paragraph_format.line_spacing = 1.15
        r = bp.add_run(b)
        r.font.size = Pt(10)

    # Save DOCX
    docx_path = os.path.abspath(output_docx)
    doc.save(docx_path)
    print(f"Word document successfully created: {docx_path}")

    # Convert to PDF using win32com if available
    pdf_path = os.path.abspath(output_pdf)
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc_com = word.Documents.Open(docx_path)
        doc_com.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
        doc_com.Close()
        word.Quit()
        print(f"PDF document successfully created: {pdf_path}")
    except Exception as e:
        print(f"Word COM PDF conversion note: {e}")
        print(f"You can open {output_docx} in Microsoft Word and click File -> Save As -> PDF.")

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--name", default="[Your Name]", help="Student Full Name")
    parser.add_argument("--reg", default="[Your Registration Number]", help="Registration Number")
    parser.add_argument("--repo", default="https://github.com/YourUsername/Autoregressive-Text-Generation-FGAI", help="GitHub repo URL")
    args = parser.parse_args()

    create_submission_doc(
        student_name=args.name,
        reg_number=args.reg,
        repo_url=args.repo
    )
